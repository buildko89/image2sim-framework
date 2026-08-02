from __future__ import annotations

import argparse
import json
import re
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


REPO = Path(__file__).resolve().parents[2]
DEFAULT_INPUT = REPO / "input" / "raw_photos" / "drone3" / "物流用ドローン部品リスト_AI用.xlsx"
DEFAULT_OUTPUT = REPO / "output" / "drone3_parametric"


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="drone3部品表を正規化し、モデル化用の整理資料を作成します。")
    parser.add_argument("--input", default=str(DEFAULT_INPUT))
    parser.add_argument("--output", default=str(DEFAULT_OUTPUT))
    return parser.parse_args()


def clean_text(value: Any) -> str | None:
    if value is None:
        return None
    text = re.sub(r"\s+", " ", str(value)).strip().strip('"”“')
    return text or None


def model_quantity(item: dict[str, Any]) -> tuple[int | None, str | None]:
    part = item["part_name"] or ""
    middle = item["middle_category"] or ""
    major = item["major_category"] or ""
    if major == "機体構造" and re.search(r"(?:HFS|HFSL)6-30(?:30|60)-\d+", part):
        return 1, "部品表で取付場所ごとに1行"
    if major == "パワープラント" and middle == "モーター":
        return 8, "PDF上面図の8ローター"
    if major == "パワープラント" and middle == "ESC":
        return 8, "1モーターにつき1 ESC"
    if major == "パワープラント" and middle == "プロペラ":
        return 4, "2本/ペアを8ローターへ使用"
    if middle == "フライトコンピュータ":
        return 1, "部品表の機体側構成"
    if major == "電源" and middle == "バッテリ":
        return 1, "部品表の機体側構成"
    if major == "翼" and middle == "主翼":
        return 1, "PDF上面図と部品表"
    return None, None


def modeling_role(item: dict[str, Any]) -> str:
    major = item["major_category"] or ""
    middle = item["middle_category"] or ""
    part = item["part_name"] or ""
    if major in {"機体構造", "パワープラント", "電源", "翼"}:
        return "主要外形"
    if middle in {"フライトコンピュータ", "GPS", "LiDAR", "FPVカメラ・VTX", "通信モジュール"}:
        return "代表形状"
    if re.search(r"コネクタ|ワッシャ|ボルト|ナット|ケーブル", middle + part):
        return "メタデータのみ"
    return "簡略形状"


def extract_parts(path: Path) -> list[dict[str, Any]]:
    workbook = load_workbook(path, data_only=False, read_only=False)
    sheet = workbook["部品表"]
    current_major: str | None = None
    current_middle: str | None = None
    items: list[dict[str, Any]] = []
    for row_index in range(6, sheet.max_row + 1):
        major = clean_text(sheet.cell(row_index, 1).value)
        middle = clean_text(sheet.cell(row_index, 2).value)
        if major:
            current_major = major
        if middle:
            current_middle = middle
        maker = clean_text(sheet.cell(row_index, 3).value)
        part_name = clean_text(sheet.cell(row_index, 4).value)
        location = clean_text(sheet.cell(row_index, 5).value)
        mass_value = sheet.cell(row_index, 6).value
        mass = float(mass_value) if isinstance(mass_value, (int, float)) else None
        if not any((maker, part_name, location, mass)):
            continue
        item = {
            "source_row": row_index,
            "major_category": current_major,
            "middle_category": current_middle,
            "manufacturer_or_supplier": maker,
            "part_name": part_name,
            "location": location,
            "unit_mass_kg": mass,
        }
        quantity, basis = model_quantity(item)
        item["modeled_quantity"] = quantity
        item["quantity_basis"] = basis
        item["modeled_mass_kg"] = round(mass * quantity, 6) if mass is not None and quantity is not None else None
        item["modeling_role"] = modeling_role(item)
        items.append(item)
    return items


def build_summary(items: list[dict[str, Any]], source: Path) -> dict[str, Any]:
    category_counts = Counter(item["major_category"] for item in items)
    known_mass_by_category: dict[str, float] = defaultdict(float)
    for item in items:
        if item["modeled_mass_kg"] is not None:
            known_mass_by_category[item["major_category"]] += item["modeled_mass_kg"]
    return {
        "schema_version": "1.0",
        "source_workbook": str(source),
        "sheet": "部品表",
        "part_entry_count": len(items),
        "entry_count_by_major_category": dict(category_counts),
        "known_minimum_mass_kg": round(sum(known_mass_by_category.values()), 6),
        "known_mass_by_category_kg": {key: round(value, 6) for key, value in known_mass_by_category.items()},
        "mass_scope_note": "数量を写真・PDF・行別取付場所から確定できた項目だけの下限。ブラケット、締結部品、未記載電装の数量・質量は含まない。",
        "modeled_layout": {
            "rotor_count": 8,
            "motor_model": "T-MOTOR MN1018 KV72",
            "motor_size_mm": {"diameter": 121.0, "height": 46.5},
            "esc_model": "T-MOTOR ALPHA 60A 24S FOC ESC",
            "esc_size_mm": {"length": 115.0, "width": 57.0, "height": 26.5},
            "esc_size_source": "provisional_same_class_envelope",
            "propeller_model": "T-MOTOR G34x11.5",
            "propeller_diameter_mm": 863.6,
            "battery": "GREPOW LIDISS PL-7874172 44.4V 16Ah 710Wh",
            "frame_motor_center_span_mm": {"x": 2560.0, "y": 2445.0},
            "cross_frame_inset_from_longitudinal_motor_mm": 500.0,
            "center_frame_width_mm": 880.0,
            "landing_leg_height_mm": 500.0,
            "skid_length_mm": 1500.0,
            "wing_mm": {"span": 2500.0, "chord": 400.0},
            "wing_side_outer_motor_mount": "under_frame",
            "gps_mast_height_mm": 500.0,
        },
        "supplemental_sources": [
            "https://store.tmotor.com/product/mn1018-motor-navigator-type.html",
            "https://uav-en.tmotor.com/2020/alpha_1119/380.html",
            "https://uav-en.tmotor.com/2018/Glossy_0410/89.html",
        ],
        "items": items,
    }


def markdown_table(items: list[dict[str, Any]]) -> list[str]:
    lines = [
        "| 大区分 | 中区分 | メーカー/購入先 | 部品名 | 場所 | 単体質量kg | モデル数量 | 役割 |",
        "|---|---|---|---|---|---:|---:|---|",
    ]
    for item in items:
        cells = [
            item["major_category"] or "",
            item["middle_category"] or "",
            item["manufacturer_or_supplier"] or "",
            item["part_name"] or "",
            item["location"] or "",
            f"{item['unit_mass_kg']:.6g}" if item["unit_mass_kg"] is not None else "",
            str(item["modeled_quantity"]) if item["modeled_quantity"] is not None else "未確定",
            item["modeling_role"],
        ]
        lines.append("| " + " | ".join(cell.replace("|", "/") for cell in cells) + " |")
    return lines


def write_markdown(summary: dict[str, Any], output_path: Path) -> None:
    category_lines = [
        f"- {category}: {count}項目"
        for category, count in summary["entry_count_by_major_category"].items()
    ]
    mass_lines = [
        f"- {category}: {mass:.4f} kg"
        for category, mass in summary["known_mass_by_category_kg"].items()
    ]
    layout = summary["modeled_layout"]
    lines = [
        "# drone3 部品情報整理",
        "",
        f"元資料: {summary['source_workbook']} / シート 部品表",
        "",
        "## 概要",
        "",
        f"- 有効な部品行: {summary['part_entry_count']}項目",
        *category_lines,
        f"- 数量まで確定できる項目の既知質量下限: **{summary['known_minimum_mass_kg']:.4f} kg**",
        f"- 質量範囲: {summary['mass_scope_note']}",
        "",
        "## Blenderモデルへ反映した主要情報",
        "",
        f"- 推進系: 8ローター、{layout['motor_model']} ×8、{layout['esc_model']} ×8",
        f"- プロペラ: {layout['propeller_model']}、直径{layout['propeller_diameter_mm']:.1f} mm、4ペア",
        f"- モーター中心スパン: X={layout['frame_motor_center_span_mm']['x']:.0f} mm、Y={layout['frame_motor_center_span_mm']['y']:.0f} mm",
        f"- 2560 mm横フレーム: 2445 mm材端モーター位置から{layout['cross_frame_inset_from_longitudinal_motor_mm']:.0f} mm内側",
        f"- 主翼側横フレーム端モーター: 下面取付、GPS支柱: {layout['gps_mast_height_mm']:.0f} mm＋円形アンテナ",
        f"- 中央フレーム幅: {layout['center_frame_width_mm']:.0f} mm、脚: {layout['landing_leg_height_mm']:.0f} mm、スキッド: {layout['skid_length_mm']:.0f} mm",
        f"- 主翼: 翼幅{layout['wing_mm']['span']:.0f} mm、翼弦{layout['wing_mm']['chord']:.0f} mm",
        f"- バッテリー: {layout['battery']}",
        "",
        "## 既知質量の内訳",
        "",
        *mass_lines,
        "",
        "## 整理上の注意",
        "",
        "- 元Excelには数量列がないため、フレームは取付場所ごとの行、推進系はPDF上面図の8ローターから数量を補いました。",
        "- ブラケット、ボルト、ナット、ワッシャ、コネクタ、ケーブルは種類を保持していますが数量未確定です。",
        "- モーター外形はメーカー公称値です。ESCとその他電装品は写真・同クラス品を基にした暫定外形です。",
        "",
        "## 全部品",
        "",
        *markdown_table(summary["items"]),
        "",
    ]
    output_path.write_text("\n".join(lines), encoding="utf-8")


def main() -> int:
    args = parse_args()
    source = Path(args.input).resolve()
    output_dir = Path(args.output).resolve()
    output_dir.mkdir(parents=True, exist_ok=True)
    items = extract_parts(source)
    summary = build_summary(items, source)
    (output_dir / "parts_inventory.json").write_text(
        json.dumps(summary, indent=2, ensure_ascii=False),
        encoding="utf-8",
    )
    write_markdown(summary, output_dir / "PARTS_SUMMARY.md")
    print(f"部品行: {len(items)}")
    print(f"既知質量下限: {summary['known_minimum_mass_kg']:.4f} kg")
    print(f"出力: {output_dir / 'PARTS_SUMMARY.md'}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
